"""Data import is physical-coordinate based; preview limits never change data."""

import copy
import datetime as dt
import io
import json
import math
import struct

import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from tavotto.workbench.data import (
    dataframe_from_snapshot,
    inspect_table,
    load_table,
    preview_data,
    select_data,
    snapshot_from_dataframe,
    transform_data,
    validate_schema,
)


def _all():
    return [{"start": 0, "end": 0}]


def _csv(text):
    return load_table(text.encode("utf-8-sig"), filename="data.csv")


def _snapshot(text):
    return select_data(_csv(text), _all(), _all())


def test_xlsx_ignores_styled_tail_and_preserves_physical_coordinates(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "原始数据"
    sheet["B2"] = "测量"
    sheet["D4"] = 3.5
    sheet["AN1000"].fill = PatternFill("solid", fgColor="FFFF00")
    book.create_sheet("第二张")["A1"] = 42
    path = tmp_path / "隧道.xlsx"
    book.save(path)

    table = load_table(path)
    info = inspect_table(table)
    assert (info["row_count"], info["column_count"]) == (4, 4)
    assert info["sheets"] == ["原始数据", "第二张"]
    selected = select_data(table, [{"start": 2, "end": 4}], [{"start": 2, "end": 4}])
    assert selected["values"] == [["测量", None, None], [None] * 3, [None, None, 3.5]]
    assert selected["row_sources"] == [2, 3, 4]
    assert selected["column_sources"] == [2, 3, 4]
    assert load_table(path, sheet="第二张").rows == [[42]]
    assert load_table(path.read_bytes(), filename=path.name, sheet=1).rows == [[42]]


def test_real_biff8_xls_reads_physical_cells_without_excel_or_a_writer_dependency():
    # This small BIFF8 stream is a real legacy Excel workbook, independently
    # constructed rather than using the parser to manufacture its own fixture.
    def record(code, payload=b""):
        return struct.pack("<HH", code, len(payload)) + payload

    def bof(kind):
        return record(0x809, struct.pack("<HHHHII", 0x600, kind, 0x0DBB, 1996, 0, 6))

    name = b"Measurements"

    def boundsheet(offset):
        return record(0x85, struct.pack("<IBBBB", offset, 0, 0, len(name), 0) + name)

    prefix = bof(5) + record(0x42, struct.pack("<H", 1200))
    offset = len(prefix) + len(boundsheet(0)) + len(record(0xA))
    worksheet = (
        bof(0x10)
        + record(0x200, struct.pack("<IIHHH", 0, 1000, 0, 40, 0))
        + record(0x204, struct.pack("<HHHHB", 1, 1, 0, 4, 0) + b"name")
        + record(0x203, struct.pack("<HHHd", 2, 2, 0, 42.5))
        + record(0x205, struct.pack("<HHHBB", 2, 1, 0, 1, 0))
        + record(0xA)
    )
    table = load_table(prefix + boundsheet(offset) + record(0xA) + worksheet, filename="legacy.xls")
    assert table.rows == [[None, None, None], [None, "name", None], [None, True, 42.5]]
    assert table.source["sheets"] == ["Measurements"]
    assert (table.row_count, table.column_count) == (3, 3)


def test_csv_quoted_newlines_blanks_and_short_rows_are_records():
    table = _csv('name,value,note\r\n"a\nb",2,""\r\n\r\nc,3\r\n')
    assert table.rows == [["name", "value", "note"], ["a\nb", 2, ""], [None] * 3, ["c", 3, None]]
    assert (table.row_count, table.column_count) == (4, 3)
    snapshot = select_data(table, _all(), _all(), header=True)
    assert snapshot["columns"] == ["name", "value", "note"]
    assert snapshot["row_sources"] == [2, 3, 4]


def test_csv_decode_is_strict_and_explicit_encoding_and_delimiter_work():
    with pytest.raises(UnicodeDecodeError):
        load_table("测量;2".encode("gbk"), filename="data.csv")
    table = load_table("测量;2".encode("gbk"), filename="data.csv", encoding="gbk", delimiter=";")
    assert table.rows == [["测量", 2]]
    with pytest.raises(ValueError):
        load_table(b'a,"unterminated', filename="data.csv")


def test_noncontiguous_segment_order_and_cross_product_are_preserved():
    table = _csv(
        "\n".join(",".join(str(10 * row + col) for col in range(1, 6)) for row in range(1, 10))
    )
    result = select_data(
        table,
        [{"start": 1, "end": 5}, {"start": 7, "end": 9}],
        [{"start": 1, "end": 2}, {"start": 4, "end": 4}],
    )
    assert (result["row_count"], result["column_count"]) == (8, 3)
    assert result["values"][-1] == [91, 92, 94]
    reordered = select_data(
        table,
        [{"start": 7, "end": 9}, {"start": 1, "end": 2}],
        [{"start": 4, "end": 4}, {"start": 1, "end": 2}],
    )
    assert reordered["row_sources"] == [7, 8, 9, 1, 2]
    assert reordered["column_sources"] == [4, 1, 2]
    assert reordered["values"][0] == [74, 71, 72]


def test_zero_means_only_outer_boundaries():
    table = _csv("1,2,3\n4,5,6\n7,8,9")
    result = select_data(table, [{"start": 0, "end": 1}, {"start": 3, "end": 0}], _all())
    assert result["values"] == [[1, 2, 3], [7, 8, 9]]
    assert select_data(table, _all(), _all())["values"] == table.rows


@pytest.mark.parametrize(
    "ranges",
    [
        [],
        [{"start": 0, "end": 0}, {"start": 3, "end": 3}],
        [{"start": 1, "end": 0}, {"start": 3, "end": 3}],
        [{"start": 1, "end": 1}, {"start": 0, "end": 3}],
        [{"start": 1, "end": 2}, {"start": 2, "end": 3}],
        [{"start": 2, "end": 1}],
        [{"start": -1, "end": 2}],
        [{"start": 1, "end": 4}],
        [{"start": 1.0, "end": 2}],
        [{"start": True, "end": 2}],
        [{"start": "1", "end": 2}],
    ],
)
def test_invalid_ranges_are_rejected_on_both_axes(ranges):
    table = _csv("1,2,3\n4,5,6\n7,8,9")
    with pytest.raises(ValueError):
        select_data(table, ranges, _all())
    with pytest.raises(ValueError):
        select_data(table, _all(), ranges)


def test_header_normalization_is_stable_and_original_header_is_retained():
    result = select_data(_csv(",x,x,x_2,\n1,2,3,4,5"), _all(), _all(), header=True)
    assert result["columns"] == ["Column 1", "x", "x_2", "x_2_2", "Column 5"]
    assert result["source"]["header_values"] == ["", "x", "x", "x_2", ""]
    assert result["source"]["header_row"] == 1
    assert result["row_sources"] == [2]


def test_header_is_consumed_before_nonsquare_transpose_and_swaps_are_immutable():
    initial = select_data(_csv("a,b,c\n1,2,3\n4,5,6"), _all(), _all(), header=True)
    untouched = copy.deepcopy(initial)
    transposed = transform_data(initial, "transpose")
    assert transposed["values"] == [[1, 4], [2, 5], [3, 6]]
    assert transposed["columns"] == ["2", "3"]
    assert transposed["index"] == ["a", "b", "c"]
    assert transposed["row_sources"] == [1, 2, 3]
    assert transposed["column_sources"] == [2, 3]
    assert transposed["source"]["transposed"] is True
    assert transform_data(transposed, "transpose")["values"] == initial["values"]
    assert transform_data(transposed, "transpose")["columns"] == initial["columns"]
    swapped = transform_data(initial, "swap_rows", i=1, j=2)
    assert swapped["values"] == [[4, 5, 6], [1, 2, 3]]
    assert swapped["row_sources"] == [3, 2]
    swapped = transform_data(initial, "swap_columns", i=1, j=3)
    assert swapped["values"] == [[3, 2, 1], [6, 5, 4]]
    assert swapped["columns"] == ["c", "b", "a"]
    assert swapped["column_sources"] == [3, 2, 1]
    assert initial == untouched
    for operation, i, j in [
        ("swap_rows", 0, 1),
        ("swap_columns", 4, 1),
        ("swap_rows", 1.0, 2),
        ("unknown", 1, 2),
    ]:
        with pytest.raises(ValueError):
            transform_data(initial, operation, i=i, j=j)


def test_typed_dataframe_survives_strict_json_roundtrip():
    frame = pd.DataFrame(
        {
            "nullable": pd.Series([1, pd.NA, 3], dtype="Int64"),
            "when": pd.to_datetime(["2026-01-01", None, "2026-03-01"], utc=True),
            "flag": pd.Series([True, pd.NA, False], dtype="boolean"),
            "mixed": pd.Series([dt.date(2026, 1, 2), None, "missing"], dtype=object),
            "float": [math.nan, math.inf, -math.inf],
            "text": pd.Series(["", pd.NA, "NA"], dtype="string"),
        }
    )
    frame.index = pd.Index(["r3", "r1", "r2"], name="measurement")
    snapshot = snapshot_from_dataframe(frame)
    encoded = json.dumps(snapshot, ensure_ascii=False, allow_nan=False)
    decoded = json.loads(encoded)
    restored = dataframe_from_snapshot(decoded)
    pd.testing.assert_frame_equal(restored, frame)
    assert restored.loc["r3", "mixed"] == dt.date(2026, 1, 2)
    assert restored.loc["r1", "mixed"] is None


def test_object_scalars_and_categories_survive_without_type_coercion():
    from decimal import Decimal

    values = [
        dt.datetime(2026, 1, 2, 3, 4, 5),
        dt.time(3, 4, 5),
        dt.timedelta(days=2, microseconds=1),
        Decimal("0.10000000000000000001"),
        2**63 + 1,
        complex(2, -3),
        b"\x00\xff",
        ("x", 3),
        pd.NA,
        pd.NaT,
        math.nan,
        None,
    ]
    frame = pd.DataFrame(
        {
            "mixed": pd.Series(values, dtype=object),
            "category": pd.Categorical(
                ["a", "b"] * 6, categories=["a", "b", "unused"], ordered=True
            ),
            "delta": pd.to_timedelta(list(range(12)), unit="ns"),
        }
    )
    restored = dataframe_from_snapshot(
        json.loads(json.dumps(snapshot_from_dataframe(frame), allow_nan=False))
    )
    pd.testing.assert_frame_equal(restored, frame)
    assert restored["mixed"][3] == Decimal("0.10000000000000000001")
    assert restored["mixed"][4] == 2**63 + 1
    assert restored["mixed"][11] is None


def test_dataframe_numeric_labels_and_empty_shapes_roundtrip():
    frames = [
        pd.DataFrame([[1, 2], [3, 4]], columns=[10, 20], index=pd.Index([3, 2], name="id")),
        pd.DataFrame({"empty": pd.Series([], dtype="Int64")}),
        pd.DataFrame(index=pd.Index([1, 2], name="id")),
    ]
    for frame in frames:
        pd.testing.assert_frame_equal(
            dataframe_from_snapshot(snapshot_from_dataframe(frame)), frame
        )


def test_truncated_preview_and_ragged_snapshot_cannot_be_rendered():
    snapshot = _snapshot("\n".join("1,2" for _ in range(51)))
    with pytest.raises(ValueError):
        dataframe_from_snapshot(preview_data(snapshot))
    snapshot["values"][0].pop()
    with pytest.raises(ValueError):
        dataframe_from_snapshot(snapshot)


def test_preview_51_by_51_is_only_a_display_limit():
    initial = _snapshot(
        "\n".join(",".join(str(row * 100 + col) for col in range(51)) for row in range(51))
    )
    before = copy.deepcopy(initial)
    preview = preview_data(initial, limit=500)
    assert len(preview["values"]) == len(preview["values"][0]) == 50
    assert len(preview["columns"]) == len(preview["index"]) == 50
    assert (preview["row_count"], preview["column_count"]) == (51, 51)
    assert (preview["preview_row_count"], preview["preview_column_count"]) == (50, 50)
    assert len(dataframe_from_snapshot(initial)) == 51
    assert initial == before


def test_schema_fixed_variable_and_numeric_errors_report_original_coordinates():
    initial = select_data(
        _csv("skip,x,y\nunused,1,10\nunused,nope,20\nunused,3,30"),
        [{"start": 2, "end": 4}],
        [{"start": 3, "end": 3}, {"start": 2, "end": 2}],
    )
    schema = {
        "rows": {"mode": "fixed", "value": 3},
        "columns": {"mode": "variable", "min": 2},
        "column_roles": [{"index": 2, "name": "X", "dtype": "numeric"}],
    }
    issues = validate_schema(initial, schema)
    assert len(issues) == 1
    assert (issues[0]["row"], issues[0]["column"]) == (2, 2)
    assert (issues[0]["source_row"], issues[0]["source_column"]) == (3, 2)
    assert issues[0]["code"] == "dtype_mismatch"
    transposed = transform_data(initial, "transpose")
    issues = validate_schema(transposed, {"column_roles": [{"index": 2, "dtype": "numeric"}]})
    assert (issues[0]["source_row"], issues[0]["source_column"]) == (3, 2)
    assert (
        validate_schema(initial, {"rows": {"mode": "fixed", "value": 4}})[0]["code"]
        == "row_count_mismatch"
    )
    assert (
        validate_schema(initial, {"columns": {"mode": "variable", "min": 3}})[0]["code"]
        == "column_count_mismatch"
    )
    assert (
        validate_schema(
            _snapshot("1,2\n3,4"),
            {
                "rows": {"mode": "variable", "min": 1},
                "column_roles": [{"index": 1, "dtype": "numeric"}],
            },
        )
        == []
    )
    assert initial["values"][1][1] == "nope"


def test_schema_text_missing_and_role_names():
    snapshot = select_data(_csv("name,value\na,1\n,2"), _all(), _all(), header=True)
    assert validate_schema(snapshot, {"column_roles": [{"name": "name", "dtype": "text"}]}) == []
    assert (
        validate_schema(snapshot, {"column_roles": [{"index": 2, "dtype": "text"}]})[0]["code"]
        == "dtype_mismatch"
    )
    assert (
        validate_schema(snapshot, {"column_roles": [{"name": "missing", "dtype": "any"}]})[0][
            "code"
        ]
        == "column_missing"
    )
    with pytest.raises(ValueError):
        validate_schema(snapshot, {"rows": {"mode": "variable", "min": -1}})
    with pytest.raises(ValueError):
        validate_schema(snapshot, {"column_roles": [{"index": 1, "dtype": "decimal-ish"}]})


def test_empty_imports_and_invalid_sheet_are_explicit():
    table = _csv("")
    assert inspect_table(table)["row_count"] == 0
    with pytest.raises(ValueError):
        select_data(table, _all(), _all())
    buffer = io.BytesIO()
    Workbook().save(buffer)
    with pytest.raises(ValueError):
        load_table(buffer.getvalue(), filename="data.xlsx", sheet="absent")
    with pytest.raises(ValueError):
        load_table(b"x", filename="data.zip")
